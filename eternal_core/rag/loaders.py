"""
Document loading for the RAG corpus.

Eternal_Master_Build_Plan.md Section 5.1: "RAG over your own docs and
notes — including the whitepapers and this build plan, so Eternal can
answer questions about its own design."

Loads every .md/.txt file under docs_dir. .docx is also supported directly
(via python-docx) for the case where a future document gets dropped in as a
real Word binary rather than pre-converted text — the whitepapers currently
in docs/ are already plain-text/markdown (see README, "why the whitepapers
are .md here, not .docx"), so this path exists for forward compatibility,
not because anything in the current corpus needs it.
"""
import os
from dataclasses import dataclass
from typing import List

SUPPORTED_EXTENSIONS = (".md", ".txt", ".docx")


@dataclass
class LoadedDocument:
    source_path: str
    title: str
    text: str


def _load_docx(path: str) -> str:
    import docx  # python-docx
    d = docx.Document(path)
    parts = []
    for para in d.paragraphs:
        if para.text.strip():
            parts.append(para.text)
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def load_documents(docs_dir: str) -> List[LoadedDocument]:
    if not os.path.isdir(docs_dir):
        raise FileNotFoundError(f"docs_dir does not exist: {docs_dir}")

    documents: List[LoadedDocument] = []
    for filename in sorted(os.listdir(docs_dir)):
        ext = os.path.splitext(filename)[1].lower()
        if ext not in SUPPORTED_EXTENSIONS:
            continue
        path = os.path.join(docs_dir, filename)
        if not os.path.isfile(path):
            continue

        if ext == ".docx":
            text = _load_docx(path)
        else:
            with open(path, "r", encoding="utf-8", errors="replace") as f:
                text = f.read()

        title = os.path.splitext(filename)[0]
        if text.strip():
            documents.append(LoadedDocument(source_path=path, title=title, text=text))

    return documents
